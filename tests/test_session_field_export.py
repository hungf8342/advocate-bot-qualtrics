from datetime import date, datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from advocate_bot_qualtrics.decision_tree.interactive_session import (
    InteractiveSessionState,
    record_node_answer,
)
from advocate_bot_qualtrics.decision_tree.session_field_export import (
    CORRECTION_CONFIDENCE_CAP,
    append_session_fields_row,
    build_field_confidence_rows,
)


def test_confirm_filing_yes_confidence():
    session = InteractiveSessionState(filing_date=date(2025, 9, 30))
    record_node_answer(session, "confirm_filing_date", "yes", 95)

    rows = {row.field_key: row for row in build_field_confidence_rows(session)}
    assert rows["filing_date"].confidence_pct == 95
    assert rows["filing_date"].corrected_from_complaint is False


def test_corrected_filing_capped_below_seventy():
    session = InteractiveSessionState(
        filing_date=date(2024, 6, 15),
        filing_date_changed=True,
    )
    record_node_answer(session, "get_filing_date", "submit", 85)

    rows = {row.field_key: row for row in build_field_confidence_rows(session)}
    assert rows["filing_date"].confidence_pct == CORRECTION_CONFIDENCE_CAP
    assert rows["filing_date"].corrected_from_complaint is True


def test_embedded_dispute_uses_different_complaint_node_confidence():
    session = InteractiveSessionState(
        filing_date=date(2024, 6, 15),
        filing_date_changed=True,
    )
    record_node_answer(session, "different_complaint_filing", "no", 60)

    rows = {row.field_key: row for row in build_field_confidence_rows(session)}
    assert rows["filing_date"].source_node_id == "different_complaint_filing"
    assert rows["filing_date"].confidence_pct == 60


def test_idk_skip_on_confirm_has_zero_confidence():
    session = InteractiveSessionState(filing_date=date(2025, 9, 30))
    record_node_answer(session, "confirm_filing_date", "yes", 0, skipped=True)

    rows = {row.field_key: row for row in build_field_confidence_rows(session)}
    assert rows["filing_date"].confidence_pct == 0


def test_og_creditor_date_from_date_node():
    session = InteractiveSessionState(last_payment_og_creditor=date(2024, 1, 4))
    record_node_answer(session, "additional_payment_OG_creditor", "yes", 90)
    record_node_answer(session, "get_last_payment_OG_creditor", "submit", 88)

    rows = {row.field_key: row for row in build_field_confidence_rows(session)}
    assert rows["last_payment_og_creditor"].confidence_pct == 88
    assert rows["last_payment_og_creditor"].value == date(2024, 1, 4)


def test_og_creditor_no_payment_exports_blank_confidence():
    session = InteractiveSessionState()
    record_node_answer(session, "additional_payment_OG_creditor", "no", 90)

    rows = {row.field_key: row for row in build_field_confidence_rows(session)}
    assert rows["last_payment_og_creditor"].value is None
    assert rows["last_payment_og_creditor"].confidence_pct is None


def test_append_session_fields_row_adds_rows(tmp_path: Path):
    path = tmp_path / "session_fields.xlsx"
    session = InteractiveSessionState(filing_date=date(2025, 9, 30))
    record_node_answer(session, "confirm_filing_date", "yes", 95)

    append_session_fields_row(
        path,
        session,
        completed_at=datetime(2025, 1, 1, tzinfo=timezone.utc),
    )
    append_session_fields_row(
        path,
        session,
        completed_at=datetime(2025, 1, 2, tzinfo=timezone.utc),
    )

    workbook = load_workbook(path)
    sheet = workbook.active
    assert sheet.max_row == 3
    assert sheet.cell(1, 1).value == "completed_at"
    assert sheet.cell(2, 2).value == "2025-09-30"
    assert sheet.cell(2, 3).value == 95
