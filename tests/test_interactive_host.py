from datetime import date
from unittest.mock import patch

import pytest

from advocate_bot_qualtrics.decision_tree.errors import ChatError
from advocate_bot_qualtrics.decision_tree.interactive_host import (
    MAX_TRANSCRIPT_TURNS,
    InteractiveChatEngine,
    build_user_payload,
    parse_submitted_date,
    render_interactive_node,
)
from advocate_bot_qualtrics.decision_tree.schemas import ChatTurnResponse
from advocate_bot_qualtrics.schemas.complaint_fact_sheet import ComplaintFactSheet


@pytest.fixture
def sample_facts() -> ComplaintFactSheet:
    return ComplaintFactSheet(
        date_complaint_filed=date(2025, 9, 30),
        date_user_failed_to_pay=date(2020, 12, 23),
        original_creditor_name="Midgard Bank",
        debt_collector_name="Southwest Collections, Inc.",
    )


def test_render_interactive_node_substitutes_dates_with_session(sample_facts):
    from advocate_bot_qualtrics.decision_tree.interactive_session import (
        init_session_from_fact_sheet,
    )

    session = init_session_from_fact_sheet(sample_facts)
    node = render_interactive_node("confirm_filing_date", sample_facts, session)
    assert "2025-09-30" in node.question


def test_render_interactive_node_annotates_party_names(sample_facts):
    from advocate_bot_qualtrics.decision_tree.interactive_session import (
        init_session_from_fact_sheet,
    )

    session = init_session_from_fact_sheet(sample_facts)
    node = render_interactive_node(
        "confirm_last_payment_complaint", sample_facts, session
    )
    assert "original creditor (Midgard Bank)" in node.question

    og_node = render_interactive_node(
        "additional_payment_OG_creditor", sample_facts, session
    )
    assert "original creditor (Midgard Bank)" in og_node.question

    dc_node = render_interactive_node(
        "additional_payment_debt_collector", sample_facts, session
    )
    assert "debt collector (Southwest Collections, Inc.)" in dc_node.question


def test_render_unknown_for_null_dates():
    facts = ComplaintFactSheet()
    from advocate_bot_qualtrics.decision_tree.interactive_session import (
        InteractiveSessionState,
    )

    node = render_interactive_node(
        "confirm_filing_date", facts, InteractiveSessionState()
    )
    assert "unknown" in node.question


def test_parse_submitted_date_formats():
    assert parse_submitted_date("2021-01-15") == date(2021, 1, 15)
    assert parse_submitted_date("01/15/2021") == date(2021, 1, 15)
    assert parse_submitted_date("January 4, 2024") == date(2024, 1, 4)
    assert parse_submitted_date("jan 4 2024") == date(2024, 1, 4)
    assert (
        parse_submitted_date(
            "I made an additional payment to the debt collector on January 4, 2024"
        )
        == date(2024, 1, 4)
    )
    assert parse_submitted_date("not a date") is None


def test_parse_submitted_date_picks_latest_when_multiple():
    assert parse_submitted_date("paid 01/01/2020 and again on 01/04/2024") == date(
        2024, 1, 4
    )


def test_build_user_payload_caps_transcript():
    transcript = [("user" if i % 2 == 0 else "assistant", f"msg{i}") for i in range(20)]
    payload = build_user_payload(transcript, "latest", max_turns=MAX_TRANSCRIPT_TURNS)
    assert "msg19" in payload
    assert "msg0" not in payload
    assert "Latest user message:" in payload


@patch("advocate_bot_qualtrics.decision_tree.interactive_host.process_chat")
def test_engine_advances_on_answer_node(mock_process_chat, sample_facts):
    mock_process_chat.return_value = ChatTurnResponse(
        user_intent="answer_node",
        assistant_reply="Great, moving on.",
        next_node_id="yes",
        answer_confidence_pct=100,
    )
    engine = InteractiveChatEngine.from_fact_sheet(sample_facts)
    engine.initial_messages()
    step = engine.submit("Yes, that is correct.")

    assert step.error is None
    assert engine.current_node_id == "confirm_last_payment_complaint"
    assert len(step.assistant_messages) == 1
    assert "last time you made a compliant payment" in step.assistant_messages[0].lower()
    assert "moving on" not in step.assistant_messages[0].lower()


@patch("advocate_bot_qualtrics.decision_tree.interactive_host.process_chat")
def test_engine_keeps_node_on_non_answer(mock_process_chat, sample_facts):
    mock_process_chat.return_value = ChatTurnResponse(
        user_intent="ask_about_complaint",
        assistant_reply="The amount sued for is $953.10.",
        next_node_id=None,
    )
    engine = InteractiveChatEngine.from_fact_sheet(sample_facts)
    engine.initial_messages()
    before = engine.current_node_id
    step = engine.submit("What amount are they suing for?")

    assert step.error is None
    assert engine.current_node_id == before
    assert len(step.assistant_messages) == 2
    assert step.assistant_messages[0] == "The amount sued for is $953.10."
    assert "2025-09-30" in step.assistant_messages[1]


@patch("advocate_bot_qualtrics.decision_tree.interactive_host.process_chat")
def test_ask_about_complaint_reasks_with_party_names(mock_process_chat, sample_facts):
    from advocate_bot_qualtrics.decision_tree.interactive_session import (
        init_session_from_fact_sheet,
    )

    mock_process_chat.return_value = ChatTurnResponse(
        user_intent="ask_about_complaint",
        assistant_reply="The original creditor is the bank you originally owed.",
        next_node_id=None,
    )
    engine = InteractiveChatEngine(
        facts=sample_facts,
        session=init_session_from_fact_sheet(sample_facts),
        current_node_id="additional_payment_OG_creditor",
    )
    engine.initial_messages()

    step = engine.submit("What does original creditor mean?")

    assert step.error is None
    assert engine.current_node_id == "additional_payment_OG_creditor"
    assert "original creditor (Midgard Bank)" in step.assistant_messages[0]
    assert "original creditor (Midgard Bank)" in step.assistant_messages[1]


@patch("advocate_bot_qualtrics.decision_tree.interactive_host.process_chat")
def test_ask_about_complaint_reasks_canonical_node_question(mock_process_chat, sample_facts):
    mock_process_chat.return_value = ChatTurnResponse(
        user_intent="ask_about_complaint",
        assistant_reply=(
            'A "complaint filed" date is when the lawsuit papers were officially '
            "filed with the court."
        ),
        next_node_id=None,
    )
    engine = InteractiveChatEngine.from_fact_sheet(sample_facts)
    engine.initial_messages()

    step = engine.submit(
        "I see September 30, 2025 — what does complaint filed mean?"
    )

    assert step.error is None
    assert engine.current_node_id == "confirm_filing_date"
    assert step.user_intent == "ask_about_complaint"
    assert len(step.assistant_messages) == 2
    assert "complaint filed" in step.assistant_messages[0].lower()
    assert "2025-09-30" in step.assistant_messages[1]


@patch("advocate_bot_qualtrics.decision_tree.interactive_host.process_chat")
def test_compound_unrelated_legal_advice_advances(mock_process_chat, sample_facts):
    mock_process_chat.return_value = ChatTurnResponse(
        user_intent="answer_node",
        assistant_reply="Got it.",
        next_node_id="no",
        answer_confidence_pct=95,
    )
    engine = InteractiveChatEngine.from_fact_sheet(sample_facts)
    engine.current_node_id = "threatening_arrest"
    engine.initial_messages()

    step = engine.submit(
        "No, nothing like that — but can they actually arrest me for not paying?"
    )

    assert step.error is None
    assert engine.current_node_id == "contact_third_parties"
    assert step.user_intent == "answer_node"


@patch("advocate_bot_qualtrics.decision_tree.interactive_host.process_chat")
def test_engine_api_error(mock_process_chat, sample_facts):
    mock_process_chat.side_effect = RuntimeError("network down")
    engine = InteractiveChatEngine.from_fact_sheet(sample_facts)
    engine.initial_messages()
    step = engine.submit("Yes")

    assert step.error is not None
    assert "API error" in step.error
    assert engine.current_node_id == "confirm_filing_date"


@patch("advocate_bot_qualtrics.decision_tree.interactive_host._process_terminal_qa")
def test_terminal_qa_stays_on_review_questions(mock_terminal_qa, sample_facts):
    mock_terminal_qa.return_value = ChatTurnResponse(
        user_intent="ask_about_complaint",
        assistant_reply="The complaint was filed on 2025-09-30.",
        next_node_id=None,
    )
    engine = InteractiveChatEngine.from_fact_sheet(sample_facts)
    engine.current_node_id = "review_questions"
    engine.tree_complete = True

    step = engine.submit("When was the complaint filed?")

    assert step.error is None
    assert engine.current_node_id == "review_questions"
    assert step.tree_complete is True


@patch("advocate_bot_qualtrics.decision_tree.interactive_host._process_terminal_qa")
def test_terminal_qa_api_error(mock_terminal_qa, sample_facts):
    mock_terminal_qa.side_effect = ChatError("bad turn")
    engine = InteractiveChatEngine.from_fact_sheet(sample_facts)
    engine.current_node_id = "review_questions"

    step = engine.submit("Any other questions?")

    assert step.error is not None
    assert engine.current_node_id == "review_questions"


@patch("advocate_bot_qualtrics.decision_tree.interactive_host.process_chat")
def test_hook_drain_after_sol_path(mock_process_chat, sample_facts):
    responses = [
        ChatTurnResponse(
            user_intent="answer_node",
            assistant_reply="ok",
            next_node_id="yes",
            answer_confidence_pct=100,
        ),
        ChatTurnResponse(
            user_intent="answer_node",
            assistant_reply="ok",
            next_node_id="yes",
            answer_confidence_pct=100,
        ),
        ChatTurnResponse(
            user_intent="answer_node",
            assistant_reply="ok",
            next_node_id="no",
            answer_confidence_pct=100,
        ),
        ChatTurnResponse(
            user_intent="answer_node",
            assistant_reply="ok",
            next_node_id="no",
            answer_confidence_pct=100,
        ),
    ]
    mock_process_chat.side_effect = responses

    engine = InteractiveChatEngine.from_fact_sheet(sample_facts)
    engine.initial_messages()
    engine.submit("yes filing")
    engine.submit("yes payment")
    engine.submit("no og")
    step = engine.submit("no collector")

    assert engine.current_node_id == "threatening_arrest"
    assert engine.last_outcomes.get("sol")
    assert any("SOL" in msg for msg in step.assistant_messages)


@patch("advocate_bot_qualtrics.decision_tree.interactive_host.process_chat")
def test_submit_with_unparseable_date_stays_on_node(mock_process_chat, sample_facts):
    mock_process_chat.return_value = ChatTurnResponse(
        user_intent="answer_node",
        assistant_reply="Thanks.",
        next_node_id="submit",
        answer_confidence_pct=100,
    )
    engine = InteractiveChatEngine.from_fact_sheet(sample_facts)
    engine.current_node_id = "get_last_payment_debt_collector"
    engine.initial_messages()
    before = engine.session.last_payment_complaint

    step = engine.submit("sometime last year")

    assert engine.current_node_id == "get_last_payment_debt_collector"
    assert engine.session.last_payment_complaint == before
    assert len(step.assistant_messages) == 1
    assert "couldn't read a date" in step.assistant_messages[0].lower()


@patch("advocate_bot_qualtrics.decision_tree.interactive_host.process_chat")
def test_different_complaint_yes_restarts_at_filing_confirm(mock_process_chat, sample_facts):
    mock_process_chat.return_value = ChatTurnResponse(
        user_intent="answer_node",
        assistant_reply="Ok.",
        next_node_id="yes",
        answer_confidence_pct=100,
    )
    engine = InteractiveChatEngine.from_fact_sheet(sample_facts)
    engine.current_node_id = "different_complaint_filing"
    engine.initial_messages()

    step = engine.submit("Yes, I think this is a different complaint.")

    assert engine.current_node_id == "confirm_filing_date"
    assert any("confirm the filing date again" in m.lower() for m in step.assistant_messages)


@patch("advocate_bot_qualtrics.decision_tree.interactive_host.process_chat")
def test_confirm_filing_no_routes_to_different_complaint_gate(mock_process_chat, sample_facts):
    mock_process_chat.return_value = ChatTurnResponse(
        user_intent="answer_node",
        assistant_reply="Ok.",
        next_node_id="no",
        answer_confidence_pct=100,
    )
    engine = InteractiveChatEngine.from_fact_sheet(sample_facts)
    engine.initial_messages()
    step = engine.submit("No, that filing date is wrong.")

    assert engine.current_node_id == "different_complaint_filing"
    assert step.error is None


@patch("advocate_bot_qualtrics.decision_tree.interactive_host.process_chat")
def test_same_complaint_reaches_filing_date_node_despite_seeded_date(
    mock_process_chat, sample_facts
):
    """Seeded filing_date must not auto-skip the dispute correction node."""
    mock_process_chat.return_value = ChatTurnResponse(
        user_intent="answer_node",
        assistant_reply="Ok.",
        next_node_id="no",
        answer_confidence_pct=100,
    )
    engine = InteractiveChatEngine.from_fact_sheet(sample_facts)
    engine.current_node_id = "different_complaint_filing"
    engine.initial_messages()

    step = engine.submit("No, same complaint — the extracted filing date is wrong.")

    assert engine.current_node_id == "get_filing_date"
    assert engine.session.filing_date == date(2025, 9, 30)
    assert engine.session.filing_date_changed is False
    assert step.error is None


@patch("advocate_bot_qualtrics.decision_tree.interactive_host.process_chat")
def test_filing_date_correction_updates_session(mock_process_chat, sample_facts):
    mock_process_chat.return_value = ChatTurnResponse(
        user_intent="answer_node",
        assistant_reply="Thanks.",
        next_node_id="submit",
        answer_confidence_pct=100,
    )
    engine = InteractiveChatEngine.from_fact_sheet(sample_facts)
    engine.current_node_id = "get_filing_date"
    engine.initial_messages()

    step = engine.submit("2024-06-15")

    assert engine.session.filing_date == date(2024, 6, 15)
    assert engine.session.filing_date_changed is True
    assert engine.current_node_id == "confirm_last_payment_complaint"
    assert step.error is None


@patch("advocate_bot_qualtrics.decision_tree.interactive_host.process_chat")
def test_same_complaint_reaches_last_payment_date_node_despite_seeded_date(
    mock_process_chat, sample_facts
):
    mock_process_chat.return_value = ChatTurnResponse(
        user_intent="answer_node",
        assistant_reply="Ok.",
        next_node_id="no",
        answer_confidence_pct=100,
    )
    engine = InteractiveChatEngine.from_fact_sheet(sample_facts)
    engine.current_node_id = "different_complaint_last_payment"
    engine.initial_messages()

    step = engine.submit("No, same complaint — the last payment date is wrong.")

    assert engine.current_node_id == "get_last_payment_complaint"
    assert engine.session.last_payment_complaint == date(2020, 12, 23)
    assert engine.session.last_payment_date_changed is False
    assert step.error is None


@patch("advocate_bot_qualtrics.decision_tree.interactive_host.process_chat")
def test_embedded_filing_date_on_same_complaint_skips_date_node(
    mock_process_chat, sample_facts
):
    mock_process_chat.return_value = ChatTurnResponse(
        user_intent="answer_node",
        assistant_reply="Please provide the correct filing date.",
        next_node_id="no",
        answer_confidence_pct=100,
    )
    engine = InteractiveChatEngine.from_fact_sheet(sample_facts)
    engine.current_node_id = "different_complaint_filing"
    engine.initial_messages()

    step = engine.submit(
        "No, same complaint — the correct filing date is 2024-06-15."
    )

    assert engine.session.filing_date == date(2024, 6, 15)
    assert engine.session.filing_date_changed is True
    assert engine.current_node_id == "confirm_last_payment_complaint"
    assert len(step.assistant_messages) == 1
    assert "correct filing date" not in step.assistant_messages[0].lower()
    assert "last time you made a compliant payment" in step.assistant_messages[0].lower()


@patch("advocate_bot_qualtrics.decision_tree.interactive_host.process_chat")
def test_pure_idk_on_confirm_skips_with_zero_confidence(mock_process_chat, sample_facts):
    engine = InteractiveChatEngine.from_fact_sheet(sample_facts)
    engine.initial_messages()

    step = engine.submit("I don't know.")

    mock_process_chat.assert_not_called()
    assert engine.current_node_id == "confirm_last_payment_complaint"
    record = engine.session.node_answers["confirm_filing_date"]
    assert record.branch_id == "yes"
    assert record.confidence_pct == 0
    assert record.skipped is True
    assert any("move on" in m.lower() for m in step.assistant_messages)
    assert step.error is None


@patch("advocate_bot_qualtrics.decision_tree.interactive_host.process_chat")
def test_low_confidence_shows_llm_hedge_before_next_question(mock_process_chat, sample_facts):
    hedge = (
        "You didn't sound fully sure, but we'll treat that as confirming the filing date."
    )
    mock_process_chat.return_value = ChatTurnResponse(
        user_intent="answer_node",
        assistant_reply=hedge,
        next_node_id="yes",
        answer_confidence_pct=55,
    )
    engine = InteractiveChatEngine.from_fact_sheet(sample_facts)
    engine.initial_messages()

    step = engine.submit("I think yes, that is the filing date.")

    assert step.error is None
    assert engine.session.node_answers["confirm_filing_date"].confidence_pct == 55
    assert step.assistant_messages[0] == hedge
    assert len(step.assistant_messages) == 2


@patch("advocate_bot_qualtrics.decision_tree.interactive_host.process_chat")
def test_low_confidence_falls_back_when_hedge_invalid(mock_process_chat, sample_facts):
    mock_process_chat.return_value = ChatTurnResponse(
        user_intent="answer_node",
        assistant_reply="Got it. What is the next question?",
        next_node_id="yes",
        answer_confidence_pct=55,
    )
    engine = InteractiveChatEngine.from_fact_sheet(sample_facts)
    engine.initial_messages()

    step = engine.submit("I think yes.")

    assert "move forward with what you indicated" in step.assistant_messages[0].lower()


@patch("advocate_bot_qualtrics.decision_tree.interactive_host.process_chat")
def test_high_confidence_advances_without_hedge(mock_process_chat, sample_facts):
    mock_process_chat.return_value = ChatTurnResponse(
        user_intent="answer_node",
        assistant_reply="Got it.",
        next_node_id="yes",
        answer_confidence_pct=95,
    )
    engine = InteractiveChatEngine.from_fact_sheet(sample_facts)
    engine.initial_messages()

    step = engine.submit("Yes, that is correct.")

    assert len(step.assistant_messages) == 1
    assert "move forward with what you indicated" not in step.assistant_messages[0].lower()


@patch("advocate_bot_qualtrics.decision_tree.interactive_host.process_chat")
def test_yes_with_embedded_date_saves_and_skips_date_node(mock_process_chat, sample_facts):
    mock_process_chat.return_value = ChatTurnResponse(
        user_intent="answer_node",
        assistant_reply="Got it.",
        next_node_id="yes",
        answer_confidence_pct=100,
    )
    engine = InteractiveChatEngine.from_fact_sheet(sample_facts)
    engine.current_node_id = "additional_payment_debt_collector"
    engine.initial_messages()

    step = engine.submit(
        "Yes, I made an additional payment to the debt collector on January 4, 2024"
    )

    assert engine.session.last_payment_debt_collector == date(2024, 1, 4)
    assert engine.current_node_id == "sol_computation" or engine.current_node_id == (
        "threatening_arrest"
    )
    assert step.error is None


def test_set_tree_complete_exports_once(sample_facts, tmp_path, monkeypatch):
    path = tmp_path / "session_fields.xlsx"
    monkeypatch.setattr(
        "advocate_bot_qualtrics.decision_tree.interactive_host.get_session_fields_xlsx_path",
        lambda: path,
    )
    engine = InteractiveChatEngine.from_fact_sheet(sample_facts)
    engine._set_tree_complete()

    assert engine.export_done is True
    assert engine.last_export_path == str(path.resolve())
    assert path.is_file()

    engine._set_tree_complete()
    assert path.is_file()
    from openpyxl import load_workbook

    assert load_workbook(path).active.max_row == 2
