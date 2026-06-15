"""Derive per-field confidence and append session rows to Excel."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

from advocate_bot_qualtrics.decision_tree.interactive_session import (
    InteractiveSessionState,
    NodeAnswerRecord,
)

CORRECTION_CONFIDENCE_CAP = 69  # strictly below 70%

_FIELD_SPECS: tuple[tuple[str, str, tuple[str, ...], str], ...] = (
    (
        "filing_date",
        "Date claim was filed",
        (
            "get_filing_date",
            "different_complaint_filing",
            "confirm_filing_date",
        ),
        "filing_date_changed",
    ),
    (
        "last_payment_complaint",
        "Last compliant payment date",
        (
            "get_last_payment_complaint",
            "different_complaint_last_payment",
            "confirm_last_payment_complaint",
        ),
        "last_payment_date_changed",
    ),
    (
        "last_payment_og_creditor",
        "Last payment to original creditor",
        (
            "get_last_payment_OG_creditor",
            "additional_payment_OG_creditor",
        ),
        "",
    ),
    (
        "last_payment_debt_collector",
        "Last payment to debt collector",
        (
            "get_last_payment_debt_collector",
            "additional_payment_debt_collector",
        ),
        "",
    ),
)

XLSX_HEADERS: tuple[str, ...] = (
    "completed_at",
    "date_claim_filed",
    "date_claim_filed_confidence",
    "last_compliant_payment",
    "last_compliant_payment_confidence",
    "last_payment_og_creditor",
    "last_payment_og_creditor_confidence",
    "last_payment_debt_collector",
    "last_payment_debt_collector_confidence",
    "filing_corrected",
    "last_payment_corrected",
)


@dataclass(frozen=True)
class FieldConfidenceRow:
    field_key: str
    label: str
    value: date | None
    confidence_pct: int | None
    source_node_id: str | None
    corrected_from_complaint: bool


def _pick_source_node(
    session: InteractiveSessionState, node_ids: tuple[str, ...]
) -> tuple[str, NodeAnswerRecord] | None:
    for node_id in node_ids:
        record = session.node_answers.get(node_id)
        if record is not None:
            return node_id, record
    return None


def build_field_confidence_rows(
    session: InteractiveSessionState,
) -> list[FieldConfidenceRow]:
    rows: list[FieldConfidenceRow] = []
    for field_key, label, node_ids, correction_flag in _FIELD_SPECS:
        value = getattr(session, field_key)
        if value is None:
            rows.append(
                FieldConfidenceRow(
                    field_key=field_key,
                    label=label,
                    value=None,
                    confidence_pct=None,
                    source_node_id=None,
                    corrected_from_complaint=False,
                )
            )
            continue

        source = _pick_source_node(session, node_ids)
        confidence_pct = source[1].confidence_pct if source else None
        source_node_id = source[0] if source else None

        corrected = bool(correction_flag and getattr(session, correction_flag))
        if corrected and confidence_pct is not None:
            confidence_pct = min(confidence_pct, CORRECTION_CONFIDENCE_CAP)

        rows.append(
            FieldConfidenceRow(
                field_key=field_key,
                label=label,
                value=value,
                confidence_pct=confidence_pct,
                source_node_id=source_node_id,
                corrected_from_complaint=corrected,
            )
        )
    return rows


def _fmt_date(value: date | None) -> str | None:
    return value.isoformat() if value else None


def build_wide_export_row(
    session: InteractiveSessionState,
    *,
    completed_at: datetime | None = None,
) -> dict[str, Any]:
    when = completed_at or datetime.now(timezone.utc)
    field_rows = {row.field_key: row for row in build_field_confidence_rows(session)}
    filing = field_rows["filing_date"]
    last_payment = field_rows["last_payment_complaint"]
    og = field_rows["last_payment_og_creditor"]
    collector = field_rows["last_payment_debt_collector"]

    return {
        "completed_at": when.isoformat(),
        "date_claim_filed": _fmt_date(filing.value),
        "date_claim_filed_confidence": filing.confidence_pct,
        "last_compliant_payment": _fmt_date(last_payment.value),
        "last_compliant_payment_confidence": last_payment.confidence_pct,
        "last_payment_og_creditor": _fmt_date(og.value),
        "last_payment_og_creditor_confidence": og.confidence_pct,
        "last_payment_debt_collector": _fmt_date(collector.value),
        "last_payment_debt_collector_confidence": collector.confidence_pct,
        "filing_corrected": filing.corrected_from_complaint,
        "last_payment_corrected": last_payment.corrected_from_complaint,
    }


def _row_to_cells(row: dict[str, Any]) -> list[Any]:
    return [row.get(header) for header in XLSX_HEADERS]


def append_session_fields_row(
    path: Path,
    session: InteractiveSessionState,
    *,
    completed_at: datetime | None = None,
) -> Path:
    """Append one wide row for a completed conversation to the workbook."""
    from openpyxl import Workbook, load_workbook

    path = path.resolve()
    path.parent.mkdir(parents=True, exist_ok=True)
    row = build_wide_export_row(session, completed_at=completed_at)
    cells = _row_to_cells(row)

    if path.is_file():
        workbook = load_workbook(path)
        sheet = workbook.active
        if sheet.max_row == 0 or sheet.cell(1, 1).value is None:
            sheet.append(list(XLSX_HEADERS))
        sheet.append(cells)
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "session_fields"
        sheet.append(list(XLSX_HEADERS))
        sheet.append(cells)

    workbook.save(path)
    return path
