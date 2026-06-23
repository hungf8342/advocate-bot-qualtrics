"""Generic Excel row append helper."""

from __future__ import annotations

from pathlib import Path
from typing import Any, Sequence


def append_row(
    path: Path,
    headers: Sequence[str],
    row: dict[str, Any],
    *,
    sheet_title: str = "session_fields",
) -> Path:
    """Append one wide row to a workbook, creating headers when needed."""
    from openpyxl import Workbook, load_workbook

    path = path.resolve()
    path.parent.mkdir(parents=True, exist_ok=True)
    cells = [row.get(header) for header in headers]

    if path.is_file():
        workbook = load_workbook(path)
        sheet = workbook.active
        if sheet.max_row == 0 or sheet.cell(1, 1).value is None:
            sheet.append(list(headers))
        sheet.append(cells)
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = sheet_title
        sheet.append(list(headers))
        sheet.append(cells)

    workbook.save(path)
    return path
